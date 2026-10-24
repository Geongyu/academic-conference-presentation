"""Thin editable-PPTX primitives. They do not guarantee layout or source accuracy."""
from __future__ import annotations
from pathlib import Path
from typing import Any
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip('#'))


def set_run_font(run: Any, family: str, east_asian_family: str | None = None) -> None:
    """Set Latin and East Asian OOXML typefaces; does not install/embed fonts."""
    run.font.name = family
    rpr = run._r.get_or_add_rPr()
    for tag, face in [('a:ea', east_asian_family or family), ('a:cs', family)]:
        child = rpr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1])
        if child is None:
            child = OxmlElement(tag)
            rpr.append(child)
        child.set('typeface', face)


def new_deck(theme: dict) -> Any:
    prs = Presentation()
    prs.slide_width = Inches(theme['canvas']['width_in'])
    prs.slide_height = Inches(theme['canvas']['height_in'])
    return prs


def add_text(slide: Any, text: str, x: float, y: float, w: float, h: float,
             theme: dict, size: float | None = None, bold: bool = False,
             fill: str | None = None, align: Any = PP_ALIGN.LEFT,
             name: str = 'body', vertical: Any = MSO_ANCHOR.TOP) -> Any:
    if w <= 0 or h <= 0:
        raise ValueError('Text box width and height must be positive.')
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = vertical
    for index, line in enumerate(str(text).split('\n')):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(4)
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size or theme['typography']['body_pt'])
        run.font.bold = bold
        run.font.color.rgb = color(fill or theme['tokens']['text'])
        set_run_font(run, theme['typography']['family'], theme['typography']['east_asian_family'])
    return box


def add_rect(slide: Any, x: float, y: float, w: float, h: float, fill: str,
             line: str | None = None, name: str = 'shape') -> Any:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    # Override any default theme shadow; keep scientific layouts flat.
    sppr = shape._element.spPr
    if sppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst') is None:
        sppr.append(OxmlElement('a:effectLst'))
    style = shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        effect = style.find('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
        if effect is not None:
            effect.set('idx', '0')
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_slide(prs: Any, theme: dict, title: str, number: int | None = None,
              citation: str = '') -> Any:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(theme['tokens']['background'])
    width, height = theme['canvas']['width_in'], theme['canvas']['height_in']
    margin = theme['canvas']['margin_in']
    add_text(slide, title, margin, 0.40, width - 2 * margin, 1.02, theme,
             size=theme['typography']['title_pt'], bold=True, name='title')
    add_rect(slide, margin, 1.48, 0.62, 0.045, theme['tokens']['primary'], name='title-accent')
    if citation:
        add_text(slide, citation, margin, height - 0.40, width - 2 * margin - 0.70,
                 0.28, theme, size=theme['typography']['citation_pt'],
                 fill=theme['tokens']['muted_text'], name='citation')
    if number is not None:
        add_text(slide, str(number), width - margin - 0.50, height - 0.40,
                 0.50, 0.28, theme, size=12, fill=theme['tokens']['muted_text'],
                 align=PP_ALIGN.RIGHT, name='slide-number')
    return slide


def add_picture_contain(slide: Any, image: Path, x: float, y: float,
                        w: float, h: float, name: str = 'source-figure') -> Any:
    image = Path(image)
    if not image.is_file():
        raise FileNotFoundError(f'Figure not found: {image}')
    with Image.open(image) as img:
        iw, ih = img.size
    if min(iw, ih, w, h) <= 0:
        raise ValueError('Image and target dimensions must be positive.')
    scale = min(w / iw, h / ih)
    fw, fh = iw * scale, ih * scale
    shape = slide.shapes.add_picture(str(image), Inches(x + (w - fw) / 2),
                                     Inches(y + (h - fh) / 2), Inches(fw), Inches(fh))
    shape.name = name
    return shape


def set_notes(slide: Any, *, slide_id: str, title: str, objective: str,
              explanation: str, cue: str, transition: str, seconds: int | float,
              source: str, backup: bool = False) -> str:
    if seconds < 0:
        raise ValueError('Speaking-time estimate must be nonnegative.')
    notes = (f'{slide_id} | {title}\n'
             f'Type: {"backup" if backup else "main"}\n'
             f'Estimated speaking time: {seconds:g} seconds\n\n'
             f'Objective / 목적\n{objective}\n\n'
             f'Explanation / 설명\n{explanation}\n\n'
             f'Visual cue / 가리킬 부분\n{cue}\n\n'
             f'Transition / 연결\n{transition}\n\n'
             f'Source and caveats / 출처와 주의점\n{source}')
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        raise RuntimeError('No notes text placeholder. Supply a valid notes master or repair the template.')
    frame.text = notes
    return notes


def save_deck(prs: Any, destination: Path, notes: list[str]) -> None:
    destination = Path(destination)
    if destination.suffix.lower() != '.pptx':
        raise ValueError('Presentation output must end with .pptx.')
    if len(notes) != len(prs.slides):
        raise ValueError('Provide one notes record per slide, including backup slides.')
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))
    # Reopen to confirm actual packaging and embedded notes, not visual correctness.
    reopened = Presentation(str(destination))
    if len(reopened.slides) != len(notes):
        raise RuntimeError('Saved presentation slide count mismatch.')
    for index, (slide, expected) in enumerate(zip(reopened.slides, notes), 1):
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if frame is None or frame.text != expected:
            raise RuntimeError(f'Embedded notes mismatch on slide {index}.')
    destination.with_name('speaker-notes.md').write_text(
        '# Speaker notes\n\n' + '\n\n---\n\n'.join(notes) + '\n', encoding='utf-8')
