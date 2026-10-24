#!/usr/bin/env python3
"""Check PPTX structure only; visual and scientific verification remain NOT_RUN."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
from pathlib import Path
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


def unresolved_markers(text: str) -> list[str]:
    return sorted(set(re.findall(r'\[(?:NEEDS SOURCE|TBD|TODO)\]|<[^<>\n]+>', text)))


def walk_shapes(shapes, in_group: bool = False):
    for shape in shapes:
        yield shape, in_group
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes, True)


def check_deck(path: Path) -> dict:
    if not path.is_file():
        raise FileNotFoundError(f'PPTX not found: {path}')
    prs = Presentation(str(path))
    errors, warnings = [], []
    notes_count = 0
    tolerance = Inches(0.01)
    if not prs.slides:
        errors.append({'slide': None, 'issue': 'Presentation contains no slides.'})
    for number, slide in enumerate(prs.slides, 1):
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        if frame is None or not frame.text.strip():
            errors.append({'slide': number, 'issue': 'Embedded speaker notes missing or empty.'})
        else:
            notes_count += 1
            for marker in unresolved_markers(frame.text):
                errors.append({'slide': number, 'shape': 'speaker-notes', 'issue': f'Unresolved marker: {marker}'})
        for shape, in_group in walk_shapes(slide.shapes):
            record = {'slide': number, 'shape': shape.name}
            if not in_group and (shape.left < -tolerance or shape.top < -tolerance or
                shape.left + shape.width > prs.slide_width + tolerance or
                shape.top + shape.height > prs.slide_height + tolerance):
                errors.append({**record, 'issue': 'Outer shape bounds extend beyond slide.'})
            if shape.rotation:
                warnings.append({**record, 'issue': 'Rotated geometry requires rendered inspection.'})
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                warnings.append({**record, 'issue': 'Group-internal geometry requires manual inspection.'})
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if shape.has_table:
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for text_frame in frames:
                for marker in unresolved_markers(text_frame.text):
                    errors.append({**record, 'issue': f'Unresolved marker: {marker}'})
                for paragraph in text_frame.paragraphs:
                    for run in paragraph.runs:
                        if (run.text.strip() and run.font.size is not None and
                            run.font.size.pt < 18 and shape.name not in ('citation', 'slide-number')):
                            warnings.append({**record, 'issue': 'Explicit text below 18 pt; review its role.',
                                             'font_pt': run.font.size.pt, 'text': run.text[:100]})
    return {
        'pptx_file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'slide_count': len(prs.slides), 'slides_with_notes': notes_count,
        'automated_structure': 'FAIL' if errors else 'PASS',
        'errors': errors, 'warnings': warnings,
        'visual_render_inspection': 'NOT_RUN', 'scientific_verification': 'NOT_RUN',
        'limitations': ['No proof of text fit, overlaps, font availability, or glyph coverage.',
                       'No source comparison or validation of figure labels and data.',
                       'Only outer nonrotated shape boxes are checked precisely.',
                       'Masters, group geometry, and chart text require rendered/manual checks.',
                       'A PASS here is not a READY release decision.'],
    }


def main() -> int:
    from preflight import configure_console
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pptx', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    try:
        report = check_deck(args.pptx)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(f"Structure: {report['automated_structure']}; "
              f"errors={len(report['errors'])}, warnings={len(report['warnings'])}. "
              'Visual and scientific QA: NOT_RUN.')
        return 1 if report['errors'] else 0
    except Exception as exc:
        print(f'Error: could not check the presentation: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
