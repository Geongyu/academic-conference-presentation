#!/usr/bin/env python3
"""Cross-check slide plans and evidence maps; speaking times are approximate."""
from __future__ import annotations
import argparse
import json
import math
from pathlib import Path
import re
import sys
from preflight import configure_console


def speaking_seconds(text: str, language: str) -> list[float]:
    """Approximate range: Korean 300-350 syllables/min, English 130-150 words/min."""
    korean = len(re.findall(r'[가-힣]', text))
    words = len(re.findall(r"[A-Za-z0-9]+(?:['’-][A-Za-z0-9]+)*", text))
    if language.lower().startswith(('ko', '한국')):
        return [60 * (korean / 350 + words / 150), 60 * (korean / 300 + words / 130)]
    return [60 * (words / 150 + korean / 350), 60 * (words / 130 + korean / 300)]


def check_plan(plan_path: Path, evidence_path: Path, pptx: Path | None = None) -> dict:
    plan_path, evidence_path = Path(plan_path), Path(evidence_path)
    plan = json.loads(plan_path.read_text(encoding='utf-8-sig'))
    evidence = json.loads(evidence_path.read_text(encoding='utf-8-sig'))
    errors, warnings, estimates = [], [], []

    def required(record, fields, label):
        if not isinstance(record, dict):
            errors.append(f'{label}: expected object')
            return {}
        for field in fields.split():
            if field not in record or record[field] is None or record[field] == '':
                errors.append(f'{label}: missing {field}')
        return record

    def records(container, field):
        value = container.get(field)
        if not isinstance(value, list):
            errors.append(f'{field}: expected array')
            return []
        return value

    def index(rows, label, fields):
        result = {}
        for row in rows:
            row = required(row, 'id ' + fields, label)
            key = row.get('id')
            if not isinstance(key, str):
                errors.append(f'{label}: invalid ID')
                continue
            if key in result:
                errors.append(f'{label}: duplicate ID {key}')
            result[key] = row
        return result

    def local_file(value, base, label):
        if not isinstance(value, str) or not (base / value).is_file():
            errors.append(f'{label}: local file missing: {value}')

    def numeric(value, label):
        if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value) or value < 0:
            errors.append(f'{label}: expected finite nonnegative seconds')
            return 0
        return value

    if not isinstance(plan, dict) or not isinstance(evidence, dict):
        raise ValueError('Plan and evidence map must be JSON objects.')
    required(plan, 'revision theme_file available_speaking_seconds target_speaking_seconds approval slides', 'plan')
    sources = index(records(evidence, 'sources'), 'source', 'path_or_identifier version accessed source_type')
    claims = index(records(evidence, 'claims'), 'claim', 'statement source_id locator verification_status')
    assets = index(records(evidence, 'assets'), 'asset', 'source_id locator path transformation caption_checked scientific_context_preserved')
    slides = records(plan, 'slides')
    index(slides, 'slide', 'main_or_backup title purpose claim_ids evidence_locators layout visual_assets estimated_seconds notes')
    if not slides:
        errors.append('Plan contains no slides.')
    local_file(plan.get('theme_file'), plan_path.parent, 'theme_file')
    approval = required(plan.get('approval'), 'status', 'approval')
    if approval.get('status') != 'approved':
        warnings.append('Approval is not approved; do not treat this plan as released.')
    elif not approval.get('instruction_evidence') or approval.get('approved_revision') != plan.get('revision'):
        errors.append('Approved plan requires instruction_evidence and matching approved_revision.')
    for key, row in sources.items():
        identifier = row.get('path_or_identifier', '')
        if isinstance(identifier, str) and not re.match(r'^(?:https?://|doi:|10\.\d{4,9}/)', identifier, re.I):
            local_file(identifier, evidence_path.parent, key)
        if row.get('accessed') is not True:
            warnings.append(f'{key}: source not accessed')
    for collection in (claims, assets):
        for key, row in collection.items():
            ref = row.get('source_id')
            if not isinstance(ref, str) or ref not in sources:
                errors.append(f'{key}: unknown source_id {ref}')
    for key, row in claims.items():
        if row.get('verification_status') != 'verified':
            warnings.append(f'{key}: claim verification is {row.get("verification_status")}')
    for key, row in assets.items():
        local_file(row.get('path'), evidence_path.parent, key)
        if row.get('caption_checked') is not True or row.get('scientific_context_preserved') is not True:
            warnings.append(f'{key}: caption/context review incomplete')
    total = 0
    for slide in slides:
        if not isinstance(slide, dict):
            continue
        key = slide.get('id', '?')
        kind = slide.get('main_or_backup')
        if kind not in ('main', 'backup'):
            errors.append(f'{key}: main_or_backup must be main or backup')
        seconds = numeric(slide.get('estimated_seconds'), key)
        if kind == 'main':
            total += seconds
        for field, lookup in [('claim_ids', claims), ('visual_assets', assets)]:
            for ref in records(slide, field):
                if not isinstance(ref, str) or ref not in lookup:
                    errors.append(f'{key}: unknown {field} reference {ref}')
        for locator in records(slide, 'evidence_locators'):
            if not isinstance(locator, str) or locator.split(':', 1)[0].strip() not in sources:
                errors.append(f'{key}: unknown evidence source in {locator}')
        notes = required(slide.get('notes'), 'language objective explanation visual_cue transition source_and_caveats', key + '.notes')
        text = str(notes.get('explanation', '')) + ' ' + str(notes.get('transition', ''))
        interval = speaking_seconds(text, str(notes.get('language', 'ko')))
        estimates.append({'slide': key, 'approximate_seconds_range': interval})
        if seconds and (interval[0] > seconds * 1.25 or interval[1] < seconds * 0.5):
            warnings.append(f'{key}: notes length and estimated_seconds differ; approximate range {interval[0]:.1f}-{interval[1]:.1f}s')
    available = numeric(plan.get('available_speaking_seconds'), 'available_speaking_seconds')
    target = numeric(plan.get('target_speaking_seconds'), 'target_speaking_seconds')
    budget = min(target, available * 0.9)
    if total > budget:
        warnings.append(f'Main slides total {total:g}s exceeds budget {budget:g}s (90% of available, capped by target).')
    if re.search(r'\[(?:NEEDS SOURCE|TBD|TODO)\]|<[^<>\n]+>', json.dumps([plan, evidence], ensure_ascii=False)):
        errors.append('Unresolved placeholder/marker in plan or evidence map.')
    if pptx is not None:
        from pptx import Presentation
        prs = Presentation(str(pptx))
        if len(prs.slides) != len(slides):
            errors.append('PPTX/plan slide count mismatch.')
        for number, (actual, planned) in enumerate(zip(prs.slides, slides), 1):
            if not isinstance(planned, dict):
                continue
            title = actual.shapes.title
            title = title or next((s for s in actual.shapes if s.name == 'title'), None)
            if title is None or title.text != planned.get('title'):
                errors.append(f'Slide {number}: PPTX/plan title mismatch.')
            frame = actual.notes_slide.notes_text_frame if actual.has_notes_slide else None
            text = frame.text.replace('\r\n', '\n') if frame is not None else ''
            for field in ('explanation', 'transition'):
                expected = planned.get('notes', {}).get(field, '')
                if expected and expected.replace('\r\n', '\n') not in text:
                    errors.append(f'Slide {number}: PPTX notes missing planned {field}.')
    return {'status': 'FAIL' if errors else 'PASS', 'errors': errors, 'warnings': warnings,
            'main_seconds': total, 'budget_seconds': budget, 'notes_estimates': estimates,
            'timing_method': 'Approximation only: Korean 300-350 syllables/min; English 130-150 words/min. Explanation + transition; rehearse with pauses and figures.'}


def main() -> int:
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--plan', required=True, type=Path)
    parser.add_argument('--evidence', required=True, type=Path)
    parser.add_argument('--pptx', type=Path)
    parser.add_argument('--out', type=Path)
    args = parser.parse_args()
    try:
        report = check_plan(args.plan, args.evidence, args.pptx)
        text = json.dumps(report, ensure_ascii=False, indent=2)
        if args.out:
            args.out.parent.mkdir(parents=True, exist_ok=True)
            args.out.write_text(text + '\n', encoding='utf-8')
        print(text)
        return 1 if report['errors'] else 0
    except (OSError, ValueError, TypeError, KeyError) as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
