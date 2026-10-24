#!/usr/bin/env python3
"""Regression fixtures for script behavior; outputs stay in the requested directory."""
from __future__ import annotations
import argparse
import io
import json
from pathlib import Path
import subprocess
import sys
from unittest.mock import patch
from pptx import Presentation
from pptx.util import Inches
from check_deck import check_deck
from make_theme import make_theme
from plan_check import check_plan
from pptx_helpers import (new_deck, add_slide, add_title_slide, add_text, add_rich_text, style_table,
                          normalize_notes,
                          set_notes, save_deck)
from preflight import configure_console, font_inventory


def run_checks(output: Path, theme: dict | None = None, overwrite: bool = False) -> None:
    output.mkdir(parents=True, exist_ok=True)
    theme = theme or make_theme()
    with patch('preflight.subprocess.run', side_effect=ValueError('bad font output')):
        assert 'korean_candidates_found' in font_inventory()
    with patch('preflight.subprocess.run', return_value=subprocess.CompletedProcess([], 0, None, None)):
        assert 'korean_candidates_found' in font_inventory()
    stream = io.TextIOWrapper(io.BytesIO(), encoding='cp949')
    with patch('sys.stdout', stream):
        configure_console()
        print('Fig. 2 — μ ± −')
        stream.flush()
    cover_prs = new_deck(theme)
    cover = add_title_slide(cover_prs, theme, 'Cover', 'Authors', 'Affiliation', 'Venue / date')
    assert {shape.name for shape in cover.shapes} == {'title', 'subtitle', 'affiliation', 'venue'}
    assert cover.shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt == theme['typography']['cover_pt']
    assert normalize_notes('a_x000D_\nb_x000D_c\r\nd') == 'a\nb\nc\nd'
    prs = new_deck(theme)
    slide = add_slide(prs, theme, '두 줄 제목\nSecond title line')
    assert slide.shapes.title is not None and slide.content_top > 1.85
    title = slide.shapes.title
    assert title.top + title.height < Inches(slide.content_top)
    body = add_rich_text(slide, [{'text': '굵게', 'bold': True}, {'text': ' italic', 'italic': True}],
                         0.6, slide.content_top, 8, 1, theme)
    assert body.text_frame.paragraphs[0].runs[-1].font.italic
    table = slide.shapes.add_table(1, 1, Inches(0.6), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = '표 텍스트'
    style_table(table, theme)
    run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == theme['typography']['body_pt']
    assert run._r.rPr.get('lang') == theme['typography'].get('language', 'ko-KR')
    note = set_notes(slide, slide_id='S01', title='두 줄 제목\nSecond title line', objective='목적',
                     explanation='설명\r\n다음 줄', cue='표', transition='끝', seconds=30, source='자체 자료')
    path = output / 'regression.pptx'
    save_deck(prs, path, [note], overwrite=overwrite)
    assert (output / 'regression.speaker-notes.md').is_file()
    before = path.read_bytes()
    try:
        save_deck(prs, path, [note])
        raise AssertionError('Overwrite was allowed')
    except FileExistsError:
        pass
    try:
        save_deck(prs, path, ['wrong notes'], overwrite=True)
        raise AssertionError('Invalid notes were accepted')
    except RuntimeError:
        pass
    assert path.read_bytes() == before
    template = new_deck(theme, path)
    inherited = add_slide(template, theme, 'Template', layout_name='Title Only')
    assert inherited._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cSld').find(
        '{http://schemas.openxmlformats.org/presentationml/2006/main}bg') is None
    set_notes(slide, slide_id='S01', title='title', objective='objective', explanation='[TODO]',
              cue='cue', transition='transition', seconds=10, source='source', language='en')
    group = slide.shapes.add_group_shape()
    add_text(group, '<locator>', 0.6, 5.5, 3, 0.5, theme)
    table.cell(0, 0).text = '[NEEDS SOURCE]'
    marked = output / 'marked.pptx'
    prs.save(str(marked))
    errors = check_deck(marked)['errors']
    assert any(e.get('shape') == 'speaker-notes' for e in errors)
    assert any('<locator>' in e['issue'] for e in errors)
    assert any('[NEEDS SOURCE]' in e['issue'] for e in errors)
    (output / 'theme.json').write_text(json.dumps(theme), encoding='utf-8')
    (output / 'source.txt').write_text('fixture', encoding='utf-8')
    evidence = {'sources': [{'id': 'SRC01', 'path_or_identifier': 'source.txt', 'version': '1',
                            'accessed': True, 'source_type': 'user-data'}],
                'claims': [{'id': 'C01', 'statement': 'Fixture', 'source_id': 'SRC01',
                            'locator': 'line 1', 'verification_status': 'verified'}],
                'assets': []}
    plan = {'revision': 'v1', 'theme_file': 'theme.json', 'available_speaking_seconds': 100,
            'target_speaking_seconds': 90,
            'approval': {'status': 'approved', 'instruction_evidence': 'fixture authorization', 'approved_revision': 'v1'},
            'slides': [{'id': 'S01', 'main_or_backup': 'main', 'title': 'Fixture', 'purpose': 'test',
                        'claim_ids': ['C01'], 'evidence_locators': ['SRC01: line 1'], 'layout': 'Title Only',
                        'visual_assets': [], 'estimated_seconds': 95,
                        'notes': {'language': 'en', 'objective': 'Test', 'explanation': 'A fixture.',
                                  'visual_cue': 'Table', 'transition': 'End', 'source_and_caveats': 'Source'}}]}
    pp, ep = output / 'slide-plan.json', output / 'evidence-map.json'
    ep.write_text(json.dumps(evidence), encoding='utf-8')
    pp.write_text(json.dumps(plan), encoding='utf-8')
    report = check_plan(pp, ep)
    assert not report['errors'] and any('exceeds budget' in w for w in report['warnings'])
    plan['slides'][0]['claim_ids'] = ['C99']
    pp.write_text(json.dumps(plan), encoding='utf-8')
    assert any('C99' in e for e in check_plan(pp, ep)['errors'])
    plan['slides'][0]['claim_ids'] = ['C01']
    plan['slides'][0]['estimated_seconds'] = 30
    plan['slides'][0]['title'] = '두 줄 제목\nSecond title line'
    plan['slides'][0]['notes']['explanation'] = '설명\n다음 줄'
    plan['slides'][0]['notes']['transition'] = '끝'
    pp.write_text(json.dumps(plan), encoding='utf-8')
    assert not check_plan(pp, ep, path)['errors']
    (output / 'regression-report.json').write_text(json.dumps({'status': 'PASS', 'checks': [
        'font failures/None stdout', 'cp949 symbols', 'multiline title', 'title placeholder',
        'rich text/table/language', 'notes CRLF/sidecar/collision/atomic verification',
        'template background', 'notes/group/table markers', 'plan references/budget']}, indent=2), encoding='utf-8')


def main() -> int:
    configure_console()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out-dir', required=True, type=Path)
    args = parser.parse_args()
    run_checks(args.out_dir)
    print('Regression checks PASS')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
